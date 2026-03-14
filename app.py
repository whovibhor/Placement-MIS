# app.py — Main Flask application (modularised)
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file
import pandas as pd
import hashlib
import re as _re
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from mysql.connector import Error

import config
from helpers import (
    UPLOAD_FOLDER, EXPECTED_HEADERS, HEADER_TO_COL, DB_COLUMNS,
    DISPLAY_COLUMNS, EDITABLE_COLUMNS, NUMERIC_FLOAT_COLS, NUMERIC_INT_COLS,
    to_float_or_none, to_int_or_none, normalize_row, normalize_rows,
    get_connection, init_db, compute_analytics,
    get_cached_analytics, save_analytics_cache, invalidate_analytics_cache,
)
from routes_data import register_data_routes
from routes_cdm import register_cdm_routes

app = Flask(__name__)
app.secret_key = config.SECRET_KEY
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

# ── Register route modules ───────────────────────────────────────────────────
register_data_routes(app)
register_cdm_routes(app)


# ── Routes ───────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    return redirect(url_for("dashboard"))


@app.route("/data-hub", methods=["GET", "POST"])
def upload():
    if request.method == "GET":
        return render_template("data_hub.html")

    # ── POST: process uploaded file ──────────────────────────────────────
    file = request.files.get("file")
    if not file or file.filename == "":
        flash("No file selected.", "danger")
        return redirect(url_for("upload"))

    if not file.filename.lower().endswith(".xlsx"):
        flash("Only .xlsx files are accepted.", "danger")
        return redirect(url_for("upload"))

    try:
        df = pd.read_excel(file, engine="openpyxl")
    except Exception as e:
        flash(f"Could not read Excel file: {e}", "danger")
        return redirect(url_for("upload"))

    # ── Normalise column names (strip whitespace) ──────────────────────
    HEADER_ALIASES = {"0.1": "10%", "0.12": "12%"}
    normalized = []
    for h in df.columns:
        s = str(h).strip()
        if s in HEADER_ALIASES:
            s = HEADER_ALIASES[s]
        normalized.append(s)
    df.columns = normalized

    # ── Validate headers ─────────────────────────────────────────────────
    file_headers = set(df.columns)
    missing = [h for h in EXPECTED_HEADERS if h not in file_headers]
    if missing:
        flash(
            f"Invalid template. Please upload correct MIS format. "
            f"Missing columns: {missing}",
            "danger",
        )
        return redirect(url_for("upload"))

    df = df[EXPECTED_HEADERS]
    df.rename(columns=HEADER_TO_COL, inplace=True)
    df = df.where(pd.notnull(df), None)
    # Force object dtype so None stays None (not numpy.nan)
    for col in DB_COLUMNS:
        df[col] = df[col].astype(object).where(df[col].notna(), None)

    df["sr_no"] = df["sr_no"].apply(to_int_or_none)

    def to_str_or_none(v):
        if v is None:
            return None
        s = str(v).strip()
        if s.lower() == "nan" or s == "":
            return None
        return s

    for col in DB_COLUMNS:
        if col == "sr_no":
            continue
        df[col] = df[col].apply(to_str_or_none)

    # ── Fix percent decimals on upload (0.6 → 60) ───────────────────────
    for col in ["percent_10", "percent_12"]:
        def fix_pct(v):
            if v is None:
                return None
            # Strip GPA/CGPA prefix
            cleaned = _re.sub(r'^(C?GPA)\s*', '', v, flags=_re.IGNORECASE).strip()
            # Convert garbage like '__' to None
            if cleaned in ('__', '_', ''):
                return None
            try:
                num = float(cleaned)
                if 0 < num <= 1:
                    num = round(num * 100, 1)
                    cleaned = str(int(num)) if num == int(num) else str(num)
            except ValueError:
                pass
            return cleaned
        df[col] = df[col].apply(fix_pct)

    # ── Convert numeric columns to proper types ──────────────────────────
    for col in ["ctc", "percent_10", "percent_12", "graduation_ogpa"]:
        df[col] = df[col].apply(to_float_or_none)
    df["backlogs"] = df["backlogs"].apply(to_int_or_none)

    # ── Upsert into MySQL (batch) ────────────────────────────────────────
    inserted = 0
    updated = 0

    placeholders = ", ".join(["%s"] * len(DB_COLUMNS))
    cols_joined = ", ".join(DB_COLUMNS)
    update_clause = ", ".join(
        [f"{c} = VALUES({c})" for c in DB_COLUMNS if c != "reg_no"]
    )
    upsert_sql = (
        f"INSERT INTO students ({cols_joined}) VALUES ({placeholders}) "
        f"ON DUPLICATE KEY UPDATE {update_clause}"
    )

    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)

        # Fetch full existing data for accurate change detection
        cursor.execute("SELECT * FROM students")
        existing_rows = {row["reg_no"]: row for row in cursor.fetchall()}

        # Compare uploaded data against DB to find real changes
        values_list = []
        for _, row in df.iterrows():
            values = tuple(
                None if (v is not None and isinstance(v, float) and pd.isna(v)) else v
                for c in DB_COLUMNS for v in [row[c]]
            )
            values_list.append(values)
            reg = row["reg_no"]
            if reg not in existing_rows:
                inserted += 1
            else:
                # Field-by-field comparison for real updates
                db_row = existing_rows[reg]
                for col in DB_COLUMNS:
                    if col == "reg_no":
                        continue
                    new_val = row[col]
                    old_val = db_row.get(col)
                    # Normalize Decimal for comparison
                    if isinstance(old_val, Decimal):
                        old_val = float(old_val)
                    # Compare as strings to handle type mismatches
                    if str(new_val if new_val is not None else "") != str(old_val if old_val is not None else ""):
                        updated += 1
                        break

        # Always do the upsert (idempotent sync)
        cursor2 = conn.cursor()
        cursor2.executemany(upsert_sql, values_list)
        conn.commit()

        # Auto-set placed_date for newly placed students (via upload)
        now_date = datetime.now().strftime("%Y-%m-%d")
        for _, row in df.iterrows():
            reg = row["reg_no"]
            new_status = row.get("status")
            if new_status == "Placed":
                old_status = existing_rows.get(reg, {}).get("status") if reg in existing_rows else None
                if old_status != "Placed":
                    cursor2 = conn.cursor()
                    cursor2.execute(
                        "UPDATE students SET placed_date = %s WHERE reg_no = %s AND placed_date IS NULL",
                        (now_date, reg),
                    )
                    cursor2.close()
            elif reg in existing_rows and existing_rows[reg].get("status") == "Placed" and new_status != "Placed":
                # Student was un-placed, clear the date
                cursor2 = conn.cursor()
                cursor2.execute(
                    "UPDATE students SET placed_date = NULL WHERE reg_no = %s",
                    (reg,),
                )
                cursor2.close()
        conn.commit()

        cursor2 = conn.cursor()

        total = inserted + updated
        has_real_changes = (inserted + updated) > 0

        if has_real_changes:
            # ── Compute content hash ─────────────────────────────────────
            sorted_df = df.sort_values("reg_no").reset_index(drop=True)
            data_hash = hashlib.md5(
                sorted_df[DB_COLUMNS].to_csv(index=False).encode("utf-8")
            ).hexdigest()

            # ── Save version snapshot (batch) ────────────────────────────
            cursor3 = conn.cursor()
            cursor3.execute(
                "INSERT INTO upload_versions (filename, uploaded_at, total_records, inserted, updated, content_hash) "
                "VALUES (%s, %s, %s, %s, %s, %s)",
                (file.filename, datetime.now(), total, inserted, updated, data_hash),
            )
            version_id = cursor3.lastrowid

            snap_cols = ", ".join(DB_COLUMNS)
            snap_placeholders = ", ".join(["%s"] * (len(DB_COLUMNS) + 1))
            snap_sql = (
                f"INSERT INTO version_snapshots (version_id, {snap_cols}) "
                f"VALUES ({snap_placeholders})"
            )
            snap_values_list = []
            for _, row in df.iterrows():
                values = (version_id,) + tuple(
                    None if (v is not None and isinstance(v, float) and pd.isna(v)) else v
                    for c in DB_COLUMNS for v in [row[c]]
                )
                snap_values_list.append(values)
            cursor3.executemany(snap_sql, snap_values_list)

            conn.commit()
            cursor3.close()

            # Invalidate analytics cache after data change
            invalidate_analytics_cache()

            flash(
                f"Upload successful. {total} records processed, "
                f"{inserted} new, {updated} updated. Version #{version_id} created.",
                "success",
            )
        else:
            flash(
                f"No changes detected — uploaded data matches the current database. "
                f"No new version created.",
                "info",
            )

        cursor.close()
        conn.close()

    except Error as e:
        flash(f"Database error: {e}", "danger")

    return redirect(url_for("upload"))


@app.route("/upload")
def upload_legacy_redirect():
    return redirect(url_for("upload"), code=301)


@app.route("/download-template/student")
def download_student_template():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student MIS Template"
    ws.append(EXPECTED_HEADERS)

    # Optional sample row (blank values under exact headers)
    ws.append([None for _ in EXPECTED_HEADERS])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="student_mis_import_template.xlsx",
    )


@app.route("/download-template/cdm")
def download_cdm_template():
    import openpyxl

    cdm_headers = [
        "Company ID",
        "Company Name",
        "Date JD Received",
        "Data Shared(Y/N)",
        "Process Date",
        "Recieved By",
        "Course",
        "Position Offered",
        "CTC",
        "HR POC Name",
        "HR POC Designation",
        "HR POC Email",
        "HR POC Phone",
        "Process Mode(On-Campus / Virtual)",
        "Location",
        "Notes",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Recruitment Import Template"
    ws.append(cdm_headers)
    ws.append([None for _ in cdm_headers])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="cdm_import_template.xlsx",
    )


@app.route("/dashboard")
def dashboard():
    try:
        analytics = get_cached_analytics()
        if not analytics:
            conn = get_connection()
            cursor = conn.cursor(dictionary=True)
            cursor.execute("SELECT * FROM students ORDER BY sr_no")
            rows = cursor.fetchall()
            normalize_rows(rows)
            analytics = compute_analytics(rows)
            cursor.close()
            conn.close()
            save_analytics_cache(analytics)
    except Error:
        analytics = {}
    return render_template("dashboard.html", analytics=analytics)


# ── Admin routes ─────────────────────────────────────────────────────────────
@app.route("/drop-all", methods=["POST"])
def drop_all():
    """Delete all student records from the database."""
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM students")
        count = cursor.rowcount
        conn.commit()
        cursor.close()
        conn.close()
        invalidate_analytics_cache()
        flash(f"All student data deleted. {count} records removed.", "success")
    except Error as e:
        flash(f"Database error: {e}", "danger")
    return redirect(url_for("upload"))


@app.route("/drop-cdm", methods=["POST"])
def drop_cdm():
    """Delete all CDM data (companies, drives, HR, courses, rounds, linked students, edit log)."""
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM drive_students")
        cursor.execute("DELETE FROM drive_rounds")
        cursor.execute("DELETE FROM drive_courses")
        cursor.execute("DELETE FROM company_hr")
        cursor.execute("DELETE FROM cdm_edit_log")
        cursor.execute("DELETE FROM company_drives")
        cursor.execute("DELETE FROM companies")
        conn.commit()
        count_companies = cursor.rowcount
        cursor.close()
        conn.close()
        flash(f"All recruitment data deleted successfully.", "success")
    except Error as e:
        flash(f"Database error: {e}", "danger")
    return redirect(url_for("upload"))


# ── Logs / Version routes ────────────────────────────────────────────────────
@app.route("/audit")
def logs_page():
    """Show all upload versions and change logs."""
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM upload_versions ORDER BY uploaded_at DESC")
        rows = cursor.fetchall()
        cursor.close()
        conn.close()
        return render_template("audit.html", versions=rows)
    except Error as e:
        flash(f"Database error: {e}", "danger")
        return render_template("audit.html", versions=[])


@app.route("/audit/version/<int:version_id>")
def version_detail(version_id):
    """Show data from a specific upload version."""
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT v.* FROM version_snapshots v WHERE v.version_id = %s ORDER BY v.sr_no",
            (version_id,),
        )
        rows = cursor.fetchall()
        cursor.close()
        conn.close()
    except Error:
        rows = []
    return render_template("audit_version.html", version_id=version_id, snap_json=rows)


@app.route("/logs")
def logs_page_legacy_redirect():
    return redirect(url_for("logs_page"), code=301)


@app.route("/versions/<int:version_id>")
def version_detail_legacy_redirect(version_id):
    return redirect(url_for("version_detail", version_id=version_id), code=301)


@app.route("/api/version/<int:version_id>")
def api_version(version_id):
    """Return snapshot data for a specific version as JSON."""
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT v.*, uv.filename, uv.uploaded_at "
            "FROM version_snapshots v "
            "JOIN upload_versions uv ON v.version_id = uv.version_id "
            "WHERE v.version_id = %s ORDER BY v.sr_no",
            (version_id,),
        )
        rows = cursor.fetchall()
        for r in rows:
            if r.get("uploaded_at"):
                r["uploaded_at"] = r["uploaded_at"].strftime("%Y-%m-%d %H:%M:%S")
        cursor.close()
        conn.close()
        return jsonify({"data": rows})
    except Error as e:
        return jsonify({"data": [], "error": str(e)}), 500


@app.route("/delete-version/<int:version_id>", methods=["POST"])
def delete_version(version_id):
    """Delete a specific upload version and its snapshot data."""
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM version_snapshots WHERE version_id = %s", (version_id,))
        cursor.execute("DELETE FROM upload_versions WHERE version_id = %s", (version_id,))
        conn.commit()
        cursor.close()
        conn.close()
        flash(f"Version #{version_id} deleted.", "success")
    except Error as e:
        flash(f"Database error: {e}", "danger")
    return redirect(url_for("logs_page"))


# ── Main ─────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    init_db()
    app.run(debug=True, port=5000)
